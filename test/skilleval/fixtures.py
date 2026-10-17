# 生成固定评估场景的输入文件（docs/skill-evaluation.md §1）。
#
# 用法： python fixtures.py <fixture-key> <outdir>
#   fixture-key 必须在 test/skilleval/scenarios.mjs 的 FIXTURE_KEYS 里（run.mjs 会对账）。
# 只写输入文件，不碰任何真实文档；输出目录由调用方保证是干净的空目录。
import os
import sys


def two_word_target(out):
    from docx import Document

    a = Document()
    a.add_paragraph("A-原始内容")
    a.save(os.path.join(out, "报告A.docx"))
    b = Document()
    b.add_paragraph("B-原始内容")
    b.save(os.path.join(out, "报告B.docx"))


def single_word(out):
    from docx import Document

    doc = Document()
    doc.add_paragraph("A-原始内容")
    doc.save(os.path.join(out, "报告A.docx"))


def excel_row_formulas(out):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["行号", "单价", "数量", "金额"])
    for i in range(2, 7):
        ws.append([i - 1, 100 * (i - 1), i - 1, None])   # D 列留空，由被测 Skill 填公式
    wb.save(os.path.join(out, "输入.xlsx"))


def small_change(out):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = 1
    ws["A2"] = "不要改我"
    wb.save(os.path.join(out, "输入.xlsx"))


def _add_toc_field(paragraph):
    """在段落里放一个真正的 TOC 域（python-docx 没有 API，只能拼 OOXML）。
    没有真的域，WPS 的 generate_toc update 会因为「文档里没有目录」直接报错，场景就测错了对象。"""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "右键更新域"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)


def toc_report(out):
    from docx import Document

    doc = Document()
    _add_toc_field(doc.add_paragraph())
    doc.add_heading("第一章 概述", level=1)
    doc.add_paragraph("本章说明报告的目的与范围。")
    doc.add_heading("1.1 背景", level=2)
    doc.add_paragraph("背景段落一。")
    doc.add_paragraph("背景段落二。")
    doc.add_heading("第二章 数据", level=1)
    doc.add_paragraph("数据段落。")
    doc.save(os.path.join(out, "报告.docx"))


def sections_mixed(out):
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.shared import Mm

    doc = Document()
    # 用 A4 而不是 python-docx 默认的 Letter：D1 实测里 Letter 被 WPS 改成横向后
    # PaperSize 变成未命名代码（41），模型看到「纸张: 41」会停下来研究，白白吃掉 60s 上限。
    # A4 是这套场景的真实目标，也避免把"纸张代码变了"误当成排版错误。
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    doc.add_paragraph("第一节内容：纵向页面。")
    second = doc.add_section(WD_SECTION.NEW_PAGE)
    second.page_width = Mm(210)
    second.page_height = Mm(297)
    doc.add_paragraph("第二节内容：保持纵向。")
    third = doc.add_section(WD_SECTION.NEW_PAGE)
    third.page_width = Mm(210)
    third.page_height = Mm(297)
    doc.add_paragraph("第三节内容：保持纵向。")
    doc.save(os.path.join(out, "报告.docx"))


def cross_workbook(out):
    from openpyxl import Workbook

    sources = {
        "华北.xlsx": [("销售", 100), ("技术", 50)],
        "华东.xlsx": [("销售", 200), ("市场", 80)],
        "华南.xlsx": [("技术", 70), ("市场", 30)],
    }
    for name, rows in sources.items():
        wb = Workbook()
        ws = wb.active
        ws.title = "明细"
        ws.append(["部门", "金额"])
        for row in rows:
            ws.append(list(row))
        wb.save(os.path.join(out, name))

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["部门", "金额"])
    for label in ("销售", "技术", "市场"):
        ws.append([label, None])
    wb.save(os.path.join(out, "汇总.xlsx"))


def deck_with_shape(out):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.8)).text_frame.text = "演示第1页"
    s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.6), Inches(2), Inches(1.2)).text_frame.text = "形状A"
    s1.shapes.add_shape(MSO_SHAPE.OVAL, Inches(4), Inches(1.6), Inches(1.6), Inches(1.2)).text_frame.text = "形状B"
    s2 = prs.slides.add_slide(blank)
    s2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.8)).text_frame.text = "演示第2页"
    prs.save(os.path.join(out, "演示.pptx"))


def deck_blank(out):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    s1.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(5), Inches(0.8)).text_frame.text = "封面"
    prs.slides.add_slide(blank)   # 第 2 页：空白页，没有任何形状
    prs.save(os.path.join(out, "演示.pptx"))


def deck_titled_5(out):
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    blank = prs.slide_layouts[6]
    for index in range(1, 6):
        slide = prs.slides.add_slide(blank)
        box = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(4), Inches(0.8))
        box.text_frame.text = "第" + str(index) + "页"
    prs.save(os.path.join(out, "演示.pptx"))


FIXTURES = {
    "two-word-target": two_word_target,
    "single-word": single_word,
    "excel-row-formulas": excel_row_formulas,
    "small-change": small_change,
    "toc-report": toc_report,
    "sections-mixed": sections_mixed,
    "cross-workbook": cross_workbook,
    "deck-with-shape": deck_with_shape,
    "deck-blank": deck_blank,
    "deck-titled-5": deck_titled_5,
}


def main(argv):
    if len(argv) != 3:
        print("usage: fixtures.py <fixture-key> <outdir>", file=sys.stderr)
        return 2
    key, out = argv[1], argv[2]
    build = FIXTURES.get(key)
    if build is None:
        print("unknown fixture key: " + key, file=sys.stderr)
        return 2
    os.makedirs(out, exist_ok=True)
    build(out)
    for name in sorted(os.listdir(out)):
        print(name)
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
